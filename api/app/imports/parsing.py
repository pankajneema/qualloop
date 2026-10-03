"""Reading uploaded files: first sheet of an .xlsx or a UTF-8 .csv, header in row 1 (A-114).

Row numbers are the physical rows of the spreadsheet (header = 1); empty rows are skipped by the caller."""

import csv
import io
from collections.abc import Iterator, Sequence
from dataclasses import dataclass

from openpyxl import load_workbook

from app.imports.normalise import clean_cell


class ImportFileError(Exception):
    """The file cannot be read; `message` is safe to show to the user."""

    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


@dataclass(frozen=True)
class RawRow:
    row_number: int
    cells: Sequence[object]


def _decode_csv(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportFileError(
            "This CSV is not UTF-8. Save it as 'CSV UTF-8' from Excel and upload it again."
        ) from exc


def iter_rows(data: bytes, extension: str) -> Iterator[RawRow]:
    """Every physical row (empty ones included, so numbering stays physical), header first."""
    if extension == "csv":
        reader = csv.reader(io.StringIO(_decode_csv(data), newline=""))
        try:
            for number, cells in enumerate(reader, start=1):
                yield RawRow(number, cells)
        except csv.Error as exc:
            raise ImportFileError(
                "This CSV file could not be read. Check it and upload it again."
            ) from exc
        return
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        for number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
            yield RawRow(number, cells)
        workbook.close()
    except ImportFileError:
        raise
    except Exception as exc:
        raise ImportFileError(
            "This file is not a readable Excel workbook. Open it in Excel, save it as .xlsx and upload it again."
        ) from exc


def read_header(data: bytes, extension: str) -> list[str]:
    """Header texts of row 1 in file order, blanks dropped (merged or empty header cells)."""
    for row in iter_rows(data, extension):
        headers = [h for h in (clean_cell(c) for c in row.cells) if h]
        if not headers:
            raise ImportFileError("The first row of the file has no column names.")
        return headers
    raise ImportFileError("The file is empty.")


@dataclass(frozen=True)
class ParsedRow:
    row_number: int
    values: dict[str, str | None]  # target field -> cleaned cell


def parse_rows(data: bytes, extension: str, mapping: dict[str, str]) -> list[ParsedRow]:
    """Mapped rows below the header; rows with no value in any mapped column are skipped (not counted)."""
    header_index: dict[str, int] | None = None
    parsed: list[ParsedRow] = []
    for row in iter_rows(data, extension):
        if header_index is None:
            header_index = {}
            for index, cell in enumerate(row.cells):
                name = clean_cell(cell)
                if name and name not in header_index:
                    header_index[name] = index
            continue
        values: dict[str, str | None] = {}
        for target, source in mapping.items():
            position = header_index.get(source)
            raw = (
                row.cells[position] if position is not None and position < len(row.cells) else None
            )
            values[target] = clean_cell(raw)
        if any(v is not None for v in values.values()):
            parsed.append(ParsedRow(row.row_number, values))
    return parsed
